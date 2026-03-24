import asyncio
#from mcp.server.fastmcp import FastMCP
from fastmcp import FastMCP
import chromadb

from tools.llm_tool import mcp_generate_llm_response
from tools.memory import load_memory, save_memory

# --------------------------
# MCP Server
# --------------------------
mcp = FastMCP("my-mcp-server")

# --------------------------
# Chroma client + collection
# --------------------------
client = chromadb.PersistentClient(path="vdb")
collection = client.get_or_create_collection("docs")

# --------------------------
# VdB Search Tool
# --------------------------
@mcp.tool()
async def search_vector_db(query: str, top_k: int = 5):
    """
    Search local Chroma DB.
    """
    results = collection.query(
        query_texts=[query],
        n_results=top_k
    )

    return {
        "query": query,
        "results": [
            {
                "id": results["ids"][0][i],
                "text": results["documents"][0][i]
            }
            for i in range(len(results["ids"][0]))
        ]
    }

# --------------------------
# VdB List All Docs tool
# --------------------------
@mcp.tool()
async def vdb_list_documents():
    """
    List all unique document filenames stored in the VDB.
    """

    results = collection.get()

    if not results or not results.get("metadatas"):
        return {"documents": [], "count": 0}

    filenames = set()

    # metadatas is a flat list: [ {"filename": ...}, {"filename": ...}, ... ]
    for meta in results["metadatas"]:
        if isinstance(meta, dict) and "filename" in meta:
            filenames.add(meta["filename"])

    return {
        "documents": sorted(filenames),
        "count": len(filenames)
    }

# --------------------------
# VdB Delete Docs tool
# --------------------------


@mcp.tool()
async def vdb_delete_document(filename: str):
    """
    Delete all chunks belonging to the given document filename.
    """

    results = collection.get(where={"filename": filename})

    if not results or not results["ids"]:
        return {
            "status": "not_found",
            "filename": filename,
            "chunks_removed": 0
        }

    raw_ids = results["ids"]

    # FIX: handle both nested and flat ID structures
    if isinstance(raw_ids[0], list):
        flat_ids = [cid for group in raw_ids for cid in group]
    else:
        flat_ids = raw_ids

    collection.delete(ids=flat_ids)

    return {
        "status": "deleted",
        "filename": filename,
        "chunks_removed": len(flat_ids)
        }

# --------------------------
# VdB Add Docs tool
# --------------------------
@mcp.tool()
async def vdb_add_document(path: str):
    """
    Add a document or all .docx files in a folder to the VDB.
    Performs sentence-level chunking and stores filename metadata.
    """

    import os
    import glob
    import re
    from docx import Document

    # Sentence splitter
    def sentence_split(text: str):
        sentences = re.split(r'(?<=[.!?])\s+', text.strip())
        return [s.strip() for s in sentences if s.strip()]

    # Work out if path is file or directory
    if os.path.isdir(path):
        files = glob.glob(os.path.join(path, "*.docx"))
    else:
        files = [path]

    added_files = []
    added_chunks = 0

    for file in files:
        filename = os.path.basename(file)

        # Load the .docx file
        doc = Document(file)
        full_text = "\n".join(p.text for p in doc.paragraphs)

        # Chunk into sentences
        chunks = sentence_split(full_text)
        ids = [f"{filename}-chunk-{i}" for i in range(len(chunks))]
        metadatas = [{"filename": filename} for _ in chunks]

        # Add to collection
        collection.add(
            ids=ids,
            documents=chunks,
            metadatas=metadatas
        )

        added_files.append(filename)
        added_chunks += len(chunks)

    return {
        "status": "added",
        "files_added": added_files,
        "total_chunks_added": added_chunks
    }


# --------------------------
# VdB Update Docs tool
# --------------------------
@mcp.tool()
async def vdb_update_document(path: str):
    """
    Update a document in the VDB:
    1. Delete all existing chunks for this filename.
    2. Re-ingest the new version with sentence-level chunking.
    """

    import os
    import re
    from docx import Document

    filename = os.path.basename(path)

    # -------- Step 1: Get old chunks --------
    old = collection.get(where={"filename": filename})
    raw_ids = old.get("ids", [])

    # Shape-safe flattening
    if raw_ids:
        first = raw_ids[0]
        if isinstance(first, list):
            old_ids = [cid for group in raw_ids for cid in group]
        else:
            old_ids = raw_ids
    else:
        old_ids = []

    # Delete existing chunks
    if old_ids:
        collection.delete(ids=old_ids)

    # -------- Step 2: Load DOCX and chunk --------
    doc = Document(path)
    full_text = "\n".join(p.text for p in doc.paragraphs)

    def sentence_split(text: str):
        sentences = re.split(r'(?<=[.!?])\s+', text.strip())
        return [s.strip() for s in sentences if s.strip()]

    chunks = sentence_split(full_text)

    # New chunk IDs + metadata
    ids = [f"{filename}-chunk-{i}" for i in range(len(chunks))]
    metadatas = [{"filename": filename} for _ in chunks]

    # Add new chunks
    collection.add(
        ids=ids,
        documents=chunks,
        metadatas=metadatas
    )

    # -------- Step 3: Return --------
    return {
        "status": "updated",
        "filename": filename,
        "old_chunks_removed": len(old_ids),
        "new_chunks_added": len(chunks)
    }

# --------------------------
# LLM tool
# --------------------------
@mcp.tool()
async def generate_llm_response(question: str, context_items: list[str]):
    return await mcp_generate_llm_response(
        question=question,
        context_items=context_items
    )


# ============================================================
#   RAG Tool: query VDB → give context → ask LLM
# ============================================================
@mcp.tool()
async def rag_answer(query: str, top_k: int = 3):
    """
    Perform retrieval-augmented generation:
    1. Search the Chroma DB for relevant docs.
    2. Pass results to the LLM with the user's query.
    3. Return the final combined answer.
    """

    # Step 1: Retrieve from Chroma
    results = collection.query(
        query_texts=[query],
        n_results=top_k
    )

    # Extract docs
    context_items = [
        results["documents"][0][i]
        for i in range(len(results["documents"][0]))
    ]

    # Step 2: Call your LLM tool with retrieved context
    llm_answer = await mcp_generate_llm_response(
        question=query,
        context_items=context_items
    )

    # Step 3: Return full structured result
    return {
        "query": query,
        "chunks_used": context_items,
        "answer": llm_answer
    }


# ============================================================
#    Stateful RAG Tool
# Loads memory → searches VDB → LLM → updates memory
# ============================================================
@mcp.tool()
async def rag_chat_stateful(user_message: str, top_k: int = 3):
    """
    Stateful RAG chat tool with:
    - conversation memory
    - query rewriting for better retrieval
    - vector DB lookup
    - LLM response grounded in retrieved chunks
    """

    # ---- Load chat history ----
    history = load_memory()  # list[str]

    # ============================================================
    # ✅ STEP 1 — LLM Query Rewriting (Disambiguation)
    # ============================================================
    rewrite_prompt = (
        "Rewrite the USER'S message into a clear, explicit question using the "
        "conversation history. Keep ONLY the rewritten question. Do NOT answer it.\n\n"
        "Conversation history:\n"
        + "\n".join(history) +
        "\n\nUser message: " + user_message
    )

    rewritten_query = await mcp_generate_llm_response(
        question=rewrite_prompt,
        context_items=[]
    )

    if not isinstance(rewritten_query, str):
        rewritten_query = str(rewritten_query)

    rewritten_query = rewritten_query.strip()

    # ============================================================
    # ✅ STEP 2 — Use the rewritten query for retrieval
    # ============================================================
    results = collection.query(
        query_texts=[rewritten_query],
        n_results=top_k
    )

    chunks = [
        results["documents"][0][i]
        for i in range(len(results["documents"][0]))
    ]

    # ============================================================
    # ✅ STEP 3 — Build improved grounded prompt
    # ============================================================
    llm_context = [
        "Relevant retrieved knowledge (highest priority):",
        *chunks,
        "",
        "Conversation history (secondary priority):",
        *history,
        "",
        (
            "Using ONLY the retrieved knowledge above and the conversation "
            "history when relevant, answer the final user message."
        ),
        "",
        f"FINAL USER MESSAGE: {user_message}",
    ]

    # ---- LLM call ----
    llm_reply = await mcp_generate_llm_response(
        question=user_message,
        context_items=llm_context
    )

    if not isinstance(llm_reply, str):
        llm_reply = str(llm_reply)

    # ============================================================
    # ✅ STEP 4 — Update memory with speaker tags
    # ============================================================
    updated = history + [
        f"USER: {user_message}",
        f"ASSISTANT: {llm_reply}"
    ]

    save_memory(updated)

    # ============================================================
    # ✅ Return structured result
    # ============================================================
    return {
        "reply": llm_reply,
        "rewritten_query": rewritten_query,
        "chunks_used": chunks,
        "history_count": len(updated),
        "message": "Chat memory updated via rewritten query."
    }

# ============================================================
#   Clear Chat History Tool
# Or memory will grow forever!
# ============================================================
@mcp.tool()
async def reset_chat_memory():
    """
    Clears persistent chat history.
    """
    save_memory([])
    return {"message": "Chat memory cleared."}

# --------------------------
# Run server
# --------------------------
if __name__ == "__main__":
    mcp.run(transport="streamable-http", host="127.0.0.1", port=8000)